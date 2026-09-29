# -*- coding: utf-8 -*-
"""Genera la tabla de pruebas del prototipo (4 columnas) lista para pegar en el documento."""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_TITLE_BG = xl_color("#2F5597")
XL_HEAD_BG = xl_color("#4472C4")
XL_BORDER = xl_color("#BFBFBF")
XL_ALT_BG = xl_color("#F2F6FC")

thin = Side(style="thin", color=XL_BORDER)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

TITLE = "Pruebas del prototipo — Cliente de escritorio (Python + Tkinter)"
HEADERS = ["Prueba", "Descripción", "Resultado esperado", "Resultado obtenido"]

ROWS = [
    ("Registro de un usuario nuevo (POST /register)", "Código 201 y cuenta creada",
     "201; cuenta creada y cabecera Location devuelta"),
    ("Registro con un correo ya existente (POST /register)", "Código 409 (conflicto)",
     "409; \"El correo ya está registrado\""),
    ("Inicio de sesión correcto (POST /login)", "Código 200 y token de sesión",
     "200; devuelve token y fecha de expiración"),
    ("Inicio de sesión con contraseña incorrecta (POST /login)", "Código 401",
     "401; \"El correo o la contraseña son incorrectos\""),
    ("Reabrir la aplicación con la sesión guardada (GET /session)", "Código 200 (sesión válida)",
     "200; entra al panel sin pedir credenciales"),
    ("Cerrar sesión y reutilizar el token (POST /logout, GET /session)",
     "200 al cerrar y 401 al reutilizar", "200 y luego 401; el token queda revocado"),
    ("Modificar el perfil (PATCH /profile)", "Código 200", "200; solo cambian los campos enviados"),
    ("Cambiar la contraseña con la actual incorrecta (PATCH /profile)", "Código 403",
     "403; INVALID_CURRENT_PASSWORD"),
    ("Consultar el catálogo (GET /books)", "Código 200 con la lista de libros",
     "200; 10 libros con formato, categoría, precio y existencia"),
    ("Buscar por texto y rango de precio (GET /books/search)", "Código 200 con resultados filtrados",
     "200; resultados filtrados correctamente"),
    ("Ver el detalle de un libro (GET /books/{isbn})", "Código 200 con imágenes y conceptos",
     "200; incluye imágenes y conceptos asociados"),
    ("Dar de alta un libro (POST /books)", "Código 201", "201; el libro aparece en el catálogo"),
    ("Dar de alta un ISBN repetido (POST /books)", "Código 409",
     "409; \"El ISBN ya existe en la base de datos\""),
    ("Modificación parcial de un libro (PATCH /books/{isbn})", "Código 200, solo cambia lo enviado",
     "200; modified_fields: [\"stock\"], el resto intacto"),
    ("Modificación completa de un libro (PUT /books/{isbn})", "Código 200",
     "200; el recurso se reemplaza por completo"),
    ("PUT con campos faltantes (PUT /books/{isbn})", "Código 400",
     "400; \"PUT requiere la representación completa del libro\""),
    ("Eliminar un libro y comprobarlo (DELETE, GET)", "200 al eliminar y 404 al consultar",
     "200 y luego 404; el libro ya no existe"),
    ("Consultar un ISBN inexistente (GET /books/{isbn})", "Código 404", "404; \"Libro no encontrado\""),
    ("Detener el servicio de libros (GET /health)", "El semáforo pasa a rojo y la app no se cierra",
     "Sin respuesta; semáforo rojo y la aplicación siguió funcionando"),
    ("Restaurar el servicio de libros (GET /health)", "El semáforo vuelve a verde",
     "200; vuelve a verde sin reiniciar la aplicación"),
    ("Guardar la configuración y reabrir la aplicación", "El valor guardado se conserva",
     "page_size=24 se relee desde runtime/config.json"),
]

wb = Workbook()
ws = wb.active
ws.title = "Pruebas"
wb.properties.title = "Tabla de pruebas del prototipo"

# --- Anchor: row1 title, row2 header, rows 3..N data ---
last_row = 2 + len(ROWS)

# Title area
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
c = ws.cell(row=1, column=1, value=TITLE)
c.font = Font(bold=True, size=12, color="FFFFFFFF")
c.fill = PatternFill("solid", fgColor=XL_TITLE_BG)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 26

# Header row
for col, name in enumerate(HEADERS, start=1):
    h = ws.cell(row=2, column=col, value=name)
    h.font = Font(bold=True, color="FFFFFFFF")
    h.fill = PatternFill("solid", fgColor=XL_HEAD_BG)
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    h.border = border
ws.row_dimensions[2].height = 22

# Data rows
for i, (desc, esperado, obtenido) in enumerate(ROWS):
    r = 3 + i
    ws.cell(row=r, column=1, value=i + 1)
    ws.cell(row=r, column=2, value=desc)
    ws.cell(row=r, column=3, value=esperado)
    ws.cell(row=r, column=4, value=obtenido)
    for col in range(1, 5):
        cell = ws.cell(row=r, column=col)
        cell.border = border
        if col == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if i % 2 == 1:
            cell.fill = PatternFill("solid", fgColor=XL_ALT_BG)

# Column widths
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 34
ws.column_dimensions["D"].width = 46

# Usability: freeze header + filter
ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:D{last_row}"
ws.sheet_view.showGridLines = False

OUT = r"C:\VSCODEOMG\IntegracionWeb\integracion03\apps\Python_app\evidencias\Tabla_de_pruebas_prototipo.xlsx"
wb.save(OUT)
print("OK ->", OUT, "| filas:", len(ROWS))
